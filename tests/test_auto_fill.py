"""
Smart auto-fill: PDF text extraction and field mapping.

The mapping tests run against text captured from the real specimen PDFs in
`front-nextjs/public/test-fixtures/documents` rather than against invented
strings, because every bug found while building this was a bug about how those
documents are actually laid out: a header that reads "Company: X Date: Y", a
title line of "Company Profile", an IBAN that runs into the next label.

The expected values are the ones a human reads off the same document. If a
fixture is regenerated with different details these fail loudly, which is the
point — a silent extraction regression is the failure mode this feature has.
"""

from __future__ import annotations

import os
import unittest

import company_profile as cp
import extraction


# --- text captured verbatim from the specimen PDFs ------------------------
TRADE_LICENCE_TEXT = """SPECIMEN — TEST FIXTURE ONLY — NOT A VALID CERTIFICATE
Trade Licence / Incorporation Certificate
Form NH-PQF-001 — Item 62
Issuer: Department of Economy and Tourism, Dubai
Company: Apex Gulf Technical Solutions LLC Date: 28 September 2026
Licensee Apex Gulf Technical Solutions LLC
Licence number CN-1094821
Licence type Limited Liability Company
Date of incorporation 12 April 2015
Expiry date 11 April 2027
Registered address Unit 2401, Marina Gate Tower, Dubai Marina, Dubai, UAE
Status Active
Rashid Al-Farsi
Authorised signatory
28 September 2026
Date
"""

COMPANY_PROFILE_TEXT = """SPECIMEN — TEST FIXTURE ONLY — NOT A VALID CERTIFICATE
Company Profile
Form NH-PQF-001 — Item 64
Issuer: Apex Gulf Technical Solutions LLC
Company: Apex Gulf Technical Solutions LLC Date: 28 September 2026
Legal name Apex Gulf Technical Solutions LLC
Nature of business Engineering services
Scope of supply Turnkey mechanical and electrical maintenance
Employees 145
Turnover (most recent) AED 18,400,000
Geographic coverage United Arab Emirates, Kingdom of Saudi Arabia
Compliance contact compliance@apexgulf.example
Rashid Al-Farsi
Authorised signatory
"""

VAT_CERT_TEXT = """SPECIMEN — TEST FIXTURE ONLY — NOT A VALID CERTIFICATE
VAT Registration Certificate
Form NH-PQF-001 — Item 63
Issuer: Federal Tax Authority, UAE
Company: Apex Gulf Technical Solutions LLC Date: 28 September 2026
Taxable person Apex Gulf Technical Solutions LLC
TRN 100293847500003
Registration date 20 May 2015
Status Active — registered for VAT
"""

BANK_LETTER_TEXT = """SPECIMEN — TEST FIXTURE ONLY — NOT A VALID CERTIFICATE
Bank Details Confirmation Letter
Form NH-PQF-001 — Item 68
Issuer: Emirates NBD, Dubai Marina Branch, UAE
Company: Apex Gulf Technical Solutions LLC Date: 28 September 2026
Account name Apex Gulf Technical Solutions LLC
IBAN AE07 0331 2345 6789 0123 456
Bank and branch Emirates NBD, Dubai Marina Branch
Currency AED
Account type Current account
"""


def mapped(text):
    return {f.field: f for f in cp.extract_company_fields(text)}


class TestDateNormalisation(unittest.TestCase):
    def test_word_form(self):
        self.assertEqual(cp.normalise_date("12 April 2015"), "2015-04-12")

    def test_numeric_day_first(self):
        # UAE paperwork is day-first; a US month-first reading would silently
        # turn 03/04/2026 into March rather than April.
        self.assertEqual(cp.normalise_date("03/04/2026"), "2026-04-03")

    def test_iso_passes_through(self):
        self.assertEqual(cp.normalise_date("2026-04-03"), "2026-04-03")

    def test_impossible_date_rejected(self):
        self.assertIsNone(cp.normalise_date("31 April 2026"))

    def test_unparseable_is_none_not_a_guess(self):
        for junk in ("", "TBD", "soon", "n/a", "31/31/2026"):
            self.assertIsNone(cp.normalise_date(junk))


class TestIbanNormalisation(unittest.TestCase):
    def test_spaces_stripped(self):
        self.assertEqual(
            cp.normalise_iban("AE07 0331 2345 6789 0123 456"),
            "AE070331234567890123456",
        )

    def test_already_compact(self):
        self.assertEqual(
            cp.normalise_iban("AE070331234567890123456"),
            "AE070331234567890123456",
        )

    def test_rejects_non_iban(self):
        self.assertIsNone(cp.normalise_iban("Emirates NBD"))


class TestTradeLicence(unittest.TestCase):
    def setUp(self):
        self.f = mapped(TRADE_LICENCE_TEXT)

    def test_legal_name_not_truncated_by_trailing_date(self):
        # The header line is "Company: <name> Date: 28 September 2026". A
        # shape-based trim cut the name to "Apex Gulf Technical".
        self.assertEqual(
            self.f["legal_name"].value, "Apex Gulf Technical Solutions LLC"
        )

    def test_licence_number(self):
        self.assertEqual(self.f["trade_license_no"].value, "CN-1094821")

    def test_expiry_is_iso_for_a_date_input(self):
        self.assertEqual(self.f["trade_license_expiry"].value, "2027-04-11")

    def test_incorporation_date(self):
        self.assertEqual(self.f["date_of_incorporation"].value, "2015-04-12")

    def test_registered_address_full_not_truncated_at_u_ae(self):
        self.assertEqual(
            self.f["registered_address"].value,
            "Unit 2401, Marina Gate Tower, Dubai Marina, Dubai, UAE",
        )

    def test_signatory_name_and_designation_inferred(self):
        self.assertEqual(
            self.f["authorized_representative_name"].value, "Rashid Al-Farsi"
        )
        self.assertEqual(
            self.f["authorized_representative_designation"].value,
            "Authorised signatory",
        )
        self.assertEqual(
            self.f["authorized_representative_name"].confidence, cp.INFERRED
        )

    def test_year_of_commencement_left_empty(self):
        # Not stated on the document. Deriving it from the incorporation year
        # would be a different fact.
        self.assertNotIn("year_of_commencement", self.f)

    def test_every_value_carries_its_source_line(self):
        for field in self.f.values():
            self.assertTrue(field.source.strip(), f"{field.field} has no source")
            self.assertIn(field.confidence, (cp.EXACT, cp.INFERRED))


class TestCompanyProfile(unittest.TestCase):
    def setUp(self):
        self.f = mapped(COMPANY_PROFILE_TEXT)

    def test_title_line_is_not_mistaken_for_the_legal_name(self):
        # "Company Profile" is the document heading. Filling the mandatory
        # legal-name field with "Profile" is worse than leaving it blank.
        self.assertEqual(
            self.f["legal_name"].value, "Apex Gulf Technical Solutions LLC"
        )

    def test_titled_variants_are_also_rejected(self):
        # Every document that begins with the label matches it. Each of these
        # was found by a real specimen or a real fixture, not invented.
        for title in (
            "Company Profile",
            "Company Profile — Extract (partial)",
            "Company Details",
            "Company Information Sheet",
            "Company Registration Certificate",
            "Company: Summary of Particulars",
        ):
            with self.subTest(title=title):
                text = f"{title}\nLegal name Apex Gulf Technical Solutions LLC\n"
                got = mapped(text)
                if "legal_name" in got:
                    self.assertEqual(
                        got["legal_name"].value, "Apex Gulf Technical Solutions LLC"
                    )

    def test_a_real_name_starting_with_a_similar_word_still_works(self):
        # The guard keys on the first word, so make sure a legitimate name that
        # begins with a title-ish word is not thrown away along with the junk.
        text = "Legal name Information Resources Trading LLC\n"
        got = mapped(text)
        self.assertEqual(got["legal_name"].value, "Information Resources Trading LLC")

    def test_compliance_contact(self):
        self.assertEqual(
            self.f["compliance_contact"].value, "compliance@apexgulf.example"
        )

    def test_country_inferred_from_the_document(self):
        self.assertEqual(
            self.f["country_of_incorporation"].value, "United Arab Emirates"
        )
        self.assertEqual(
            self.f["country_of_incorporation"].confidence, cp.INFERRED
        )


class TestVatCertificate(unittest.TestCase):
    def setUp(self):
        self.f = mapped(VAT_CERT_TEXT)

    def test_trn(self):
        self.assertEqual(self.f["vat_registration_no"].value, "100293847500003")

    def test_vat_status_inferred_so_the_trn_field_reveals(self):
        # The TRN input is behind a conditional on vat_registration_status, so
        # a TRN without the status would look like the extraction failed.
        self.assertEqual(self.f["vat_registration_status"].value, "registered")

    def test_registration_date_is_not_reported_as_incorporation(self):
        # "Registration date 20 May 2015" here is the *tax* registration date.
        self.assertNotIn("date_of_incorporation", self.f)


class TestBankLetter(unittest.TestCase):
    def setUp(self):
        self.f = mapped(BANK_LETTER_TEXT)

    def test_iban_extracted_and_compacted(self):
        self.assertEqual(self.f["bank_iban"].value, "AE070331234567890123456")

    def test_bank_name_and_branch(self):
        self.assertEqual(
            self.f["bank_name_branch_country"].value,
            "Emirates NBD, Dubai Marina Branch",
        )

    def test_account_number_not_filled_from_the_iban(self):
        # "iban" used to be an alias of the account-number label, which copied
        # the IBAN into bank_account_number as well.
        self.assertNotIn("bank_account_number", self.f)

    def test_bank_fields_report_their_wizard_step(self):
        self.assertEqual(self.f["bank_iban"].step, 3)


class TestLabelShapes(unittest.TestCase):
    """How labels are written on real forms, as opposed to in the pattern table.

    Every case here was found by generating a "complete" test PDF and seeing a
    wrong value come back. They are grouped together because they are all the
    same subject: the label catalogue is keyed on short forms ("company",
    "account name"), and real paperwork prints long or qualified forms.
    """

    def test_slash_form_label_is_matched_whole(self):
        # "Company / legal name" — matching only "^company" left the value as
        # "/ legal name Apex Gulf Technical Solutions LLC".
        got = mapped("Company / legal name Apex Gulf Technical Solutions LLC")
        self.assertEqual(got["legal_name"].value, "Apex Gulf Technical Solutions LLC")

    def test_slash_form_bank_label_is_matched_whole(self):
        got = mapped("Bank name / branch / country Emirates NBD, Dubai Marina Branch")
        self.assertEqual(
            got["bank_name_branch_country"].value,
            "Emirates NBD, Dubai Marina Branch",
        )

    def test_longest_label_wins_over_its_own_prefix(self):
        # "year of commencement" is a prefix of "year of commencement of
        # business"; trying the short one first left "of business 2015".
        got = mapped("Year of commencement of business 2015")
        self.assertEqual(got["year_of_commencement"].value, "2015")

    def test_qualified_label_still_matches(self):
        # Anchored at ^, the bare "compliance contact" cannot match a line that
        # starts "Primary compliance contact".
        got = mapped("Primary compliance contact compliance@apexgulf.example")
        self.assertEqual(got["compliance_contact"].value, "compliance@apexgulf.example")

    def test_qualified_bank_account_labels(self):
        # Same problem: documents write "Bank account name", the table says
        # "account name", so the field was silently dropped.
        got = mapped(
            "Bank account name Apex Gulf Technical Solutions LLC\n"
            "Bank account number 1234567890"
        )
        self.assertEqual(got["bank_account_name"].value, "Apex Gulf Technical Solutions LLC")
        self.assertEqual(got["bank_account_number"].value, "1234567890")

    def test_qualifier_does_not_invent_a_match_in_a_sentence(self):
        # The qualifier set is deliberately closed. A word that is not in it must
        # not turn an ordinary sentence into a labelled value.
        prose = "as reported by the supplier address was unchanged last year"
        self.assertEqual(cp.extract_company_fields(prose), [])


class TestRobustness(unittest.TestCase):
    def test_empty_text(self):
        self.assertEqual(cp.extract_company_fields(""), [])
        self.assertEqual(cp.extract_company_fields("   \n\n  "), [])

    def test_unrelated_prose_yields_nothing_rather_than_guessing(self):
        prose = (
            "This is an invoice for consulting services rendered in March. "
            "Please remit payment within thirty days of receipt. Thank you "
            "for your business and we look forward to serving you again soon."
        )
        self.assertEqual(cp.extract_company_fields(prose), [])

    def test_results_are_in_catalogue_order(self):
        fields = cp.extract_company_fields(TRADE_LICENCE_TEXT)
        present = {f.field for f in fields}
        order = [name for name in cp.FIELD_LABELS if name in present]
        self.assertEqual([f.field for f in fields], order)

    def test_every_returned_key_is_a_real_form_field(self):
        fields = cp.extract_company_fields(TRADE_LICENCE_TEXT + BANK_LETTER_TEXT)
        for f in fields:
            self.assertIn(f.field, cp.FIELD_LABELS)


class TestExtraction(unittest.TestCase):
    def test_rejects_non_pdf_by_content(self):
        with self.assertRaises(extraction.UnsupportedUpload):
            extraction.extract_text(b"not a pdf at all", "licence.pdf")

    def test_word_file_named_pdf_is_named_as_such(self):
        # A .doc renamed .pdf must not be told "got .pdf"; the reply has to say
        # the file is actually Word and how to save it correctly.
        with self.assertRaises(extraction.UnsupportedUpload) as ctx:
            extraction.extract_text(b"PK\x03\x04 wordprocessingml", "licence.pdf")
        self.assertIn("Word document", str(ctx.exception))
        self.assertIn(".docx", str(ctx.exception))

    def test_pdf_named_docx_is_named_as_such(self):
        with self.assertRaises(extraction.UnsupportedUpload) as ctx:
            extraction.extract_text(b"%PDF-1.4 body", "licence.docx")
        self.assertIn(".pdf", str(ctx.exception))

    def test_undisguised_garbage_is_told_it_is_neither(self):
        with self.assertRaises(extraction.UnsupportedUpload) as ctx:
            extraction.extract_text(b"not a pdf at all", "licence.pdf")
        self.assertIn("not a PDF or a Word document", str(ctx.exception))

    def test_wrong_extension_message_names_the_extension(self):
        with self.assertRaises(extraction.UnsupportedUpload) as ctx:
            extraction.extract_text(b"%PDF-1.4 body", "licence.xlsx")
        self.assertIn(".xlsx", str(ctx.exception))

    def test_legacy_doc_refused_with_the_fix_in_the_message(self):
        # python-docx cannot read binary .doc, and the person sending it can fix
        # it in one click — so say so rather than just "unsupported".
        with self.assertRaises(extraction.UnsupportedUpload) as ctx:
            extraction.extract_text(b"\xd0\xcf\x11\xe0old binary doc", "licence.doc")
        self.assertIn(".docx", str(ctx.exception))
        self.assertIn("Save As", str(ctx.exception))

    def test_oversize(self):
        with self.assertRaises(extraction.UploadTooLarge):
            extraction.extract_text(
                b"%PDF-1.4" + b"0" * (extraction.MAX_UPLOAD_BYTES + 1), "big.pdf"
            )

    def test_limit_is_fifty_megabytes(self):
        self.assertEqual(extraction.MAX_UPLOAD_BYTES, 50 * 1024 * 1024)

    def test_file_just_under_the_limit_is_not_rejected_for_size(self):
        # 49 MB is allowed; the 413 is a size guard, not the only gate.
        blob = b"%PDF-1.4" + b"0" * (49 * 1024 * 1024)
        result = extraction.extract_text(blob, "big.pdf")
        # Whatever else is wrong with 49 MB of zeros, it must not be the size.
        self.assertNotIn("limit", result.reason or "")
        self.assertNotIn("50 MB", result.reason or "")

    def test_result_dict_omits_the_document_text(self):
        # The licence text must not be echoed back to the browser.
        self.assertNotIn("text", extraction.ExtractionResult(extracted=True, text="x").to_dict())

    def test_never_raises_on_unreadable_pdf(self):
        result = extraction.extract_text(b"%PDF-1.4 truncated garbage", "broken.pdf")
        self.assertFalse(result.extracted)
        self.assertTrue(result.reason)


def _make_docx(path, paragraphs, table_rows=None, trailer=True):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=0, cols=2)
        for label, value in table_rows:
            row = table.add_row()
            row.cells[0].text = label
            row.cells[1].text = value
    if trailer:
        doc.add_paragraph("SPECIMEN — TEST FIXTURE ONLY — NOT A VALID CERTIFICATE")
    doc.save(path)
    return path


class TestDocxExtraction(unittest.TestCase):
    """Word is what a notary template actually arrives as, often as a table."""

    @classmethod
    def setUpClass(cls):
        import tempfile

        cls.tmp = tempfile.mkdtemp()
        cls.path = _make_docx(
            os.path.join(cls.tmp, "licence.docx"),
            ["Trade Licence / Incorporation Certificate"],
            table_rows=[
                ("Legal name", "Apex Gulf Technical Solutions LLC"),
                ("Licence number", "CN-1094821"),
                ("Expiry date", "11 April 2027"),
                ("Date of incorporation", "12 April 2015"),
            ],
        )
        with open(cls.path, "rb") as fh:
            cls.data = fh.read()

    def test_docx_is_extracted(self):
        result = extraction.extract_text(self.data, "licence.docx")
        self.assertTrue(result.extracted, result.reason)

    def test_table_rows_feed_the_mapper(self):
        # The whole point of flattening tables: a two-column licence template is
        # the common case, and it must read as `Label value`.
        result = extraction.extract_text(self.data, "licence.docx")
        fields = {f.field: f.value for f in cp.extract_company_fields(result.text)}
        self.assertEqual(fields["legal_name"], "Apex Gulf Technical Solutions LLC")
        self.assertEqual(fields["trade_license_no"], "CN-1094821")
        self.assertEqual(fields["trade_license_expiry"], "2027-04-11")
        self.assertEqual(fields["date_of_incorporation"], "2015-04-12")

    def test_empty_docx_is_refused_not_crashed(self):
        path = _make_docx(os.path.join(self.tmp, "empty.docx"), [], trailer=False)
        with open(path, "rb") as fh:
            data = fh.read()
        result = extraction.extract_text(data, "empty.docx")
        self.assertFalse(result.extracted)
        self.assertTrue(result.reason)

    def test_zip_that_is_not_a_docx_is_refused_clearly(self):
        import zipfile
        import io

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("xl/workbook.xml", "<workbook/>")
        result = extraction.extract_text(buf.getvalue(), "book.docx")
        self.assertFalse(result.extracted)
        self.assertIn("not a Word document", result.reason)

    def test_corrupt_docx_never_raises(self):
        result = extraction.extract_text(
            b"PK\x03\x04 truncated garbage that claims to be a docx", "broken.docx"
        )
        self.assertFalse(result.extracted)
        self.assertTrue(result.reason)


if __name__ == "__main__":
    unittest.main()
